import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BLACK = RGBColor(0, 0, 0)
FONT_FAMILY = "Times New Roman"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table):
    """Sets complete, crisp, professional grid borders on tables."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="5" w:space="0" w:color="000000"/>'
        f'  <w:insideV w:val="single" w:sz="5" w:space="0" w:color="000000"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h.paragraph_format.space_before = Pt(9 if level == 1 else (6 if level == 2 else 3.5))
    h.paragraph_format.space_after = Pt(2)
    run = h.runs[0]
    run.font.color.rgb = BLACK
    run.font.name = FONT_FAMILY
    run.bold = True
    if level == 1:
        run.font.size = Pt(16) # 16pt for main headings
    elif level == 2:
        run.font.size = Pt(14) # 14pt for sub-headings
    elif level == 3:
        run.font.size = Pt(12)
    return h

def add_callout(doc, text, title="NOTE / ARCHITECTURAL PRINCIPLE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.0)
    set_cell_background(cell, "F9FAFB")
    set_cell_margins(cell, top=60, bottom=60, left=120, right=120)
    
    tcPr = cell._tc.get_or_add_tcPr()
    bdr = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(bdr)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"[{title}]\n")
    r_title.bold = True
    r_title.font.name = FONT_FAMILY
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = BLACK
    
    r_text = p.add_run(text)
    r_text.font.name = FONT_FAMILY
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = BLACK
    
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_code_block(doc, code_str, caption=None):
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(1.5)
        cr = cp.add_run(f"Listing: {caption}")
        cr.bold = True
        cr.font.name = FONT_FAMILY
        cr.font.size = Pt(10)
        cr.font.color.rgb = BLACK

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=50, bottom=50, left=100, right=100)
    
    tcPr = cell._tc.get_or_add_tcPr()
    bdr = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(bdr)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(code_str)
    r.font.name = 'Consolas'
    r.font.size = Pt(7.5)
    r.font.color.rgb = BLACK
    
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def format_table(table, col_widths, headers, data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "F3F4F6")
        set_cell_margins(hdr_cells[i], top=50, bottom=50, left=70, right=70)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = BLACK
            run.font.name = FONT_FAMILY
            run.font.size = Pt(10)
            
    for row_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        bg_hex = "FAFAFA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            row_cells[col_idx].text = str(val)
            set_cell_background(row_cells[col_idx], bg_hex)
            set_cell_margins(row_cells[col_idx], top=40, bottom=40, left=70, right=70)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = BLACK
                run.font.name = FONT_FAMILY
                if col_idx == 0:
                    run.font.bold = True
                    
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

def add_p(doc, text):
    """Clean left-aligned 11pt Times New Roman body paragraph (no awkward justification gaps)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.03
    r = p.add_run(text)
    r.font.color.rgb = BLACK
    r.font.name = FONT_FAMILY
    r.font.size = Pt(11) # 11pt for body
    return p

def add_bullet(doc, title, text):
    """Clean left-aligned 11pt Times New Roman bullet item."""
    bp = doc.add_paragraph(style='List Bullet')
    bp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    bp.paragraph_format.space_before = Pt(0.5)
    bp.paragraph_format.space_after = Pt(1)
    bp.paragraph_format.line_spacing = 1.03
    r_b = bp.add_run(title)
    r_b.bold = True
    r_b.font.color.rgb = BLACK
    r_b.font.name = FONT_FAMILY
    r_b.font.size = Pt(11)
    r_d = bp.add_run(text)
    r_d.font.color.rgb = BLACK
    r_d.font.name = FONT_FAMILY
    r_d.font.size = Pt(11)
    return bp
