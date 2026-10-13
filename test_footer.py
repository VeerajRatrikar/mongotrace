import docx
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client

doc = Document()
sect = doc.sections[0]

# Set start page number to 2
sectPr = sect._sectPr
pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="2"/>')
sectPr.append(pgNumType)

# Double border around page
xml_border = f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page"><w:top w:val="double" w:sz="12" w:space="24" w:color="000000"/><w:left w:val="double" w:sz="12" w:space="24" w:color="000000"/><w:bottom w:val="double" w:sz="12" w:space="24" w:color="000000"/><w:right w:val="double" w:sz="12" w:space="24" w:color="000000"/></w:pgBorders>'
sectPr.append(parse_xml(xml_border))

# Footer with "X | P a g e"
footer = sect.footer
f_p = footer.paragraphs[0]
f_p.text = ""
f_run1 = f_p.add_run()
fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
f_run1._r.append(fld)
f_run2 = f_p.add_run(" | P a g e")
f_run2.font.name = "Times New Roman"
f_run2.font.size = docx.shared.Pt(10)

doc.add_paragraph("First content page (labeled 2 | P a g e)")
doc.add_page_break()
doc.add_paragraph("Second content page (labeled 3 | P a g e)")

doc.save(r'C:\Users\VEERAJ\OneDrive\Desktop\NoSQL\test_footer.docx')

word = win32com.client.Dispatch('Word.Application')
d = word.Documents.Open(r'C:\Users\VEERAJ\OneDrive\Desktop\NoSQL\test_footer.docx')
d.SaveAs(r'C:\Users\VEERAJ\OneDrive\Desktop\NoSQL\test_footer.pdf', FileFormat=17)
d.Close()
word.Quit()
print("Footer & border test successful!")
